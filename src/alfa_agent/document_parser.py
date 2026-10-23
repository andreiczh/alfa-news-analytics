"""Bounded local parsing. Never execute macros, formulas or document instructions."""

from __future__ import annotations

import csv
import io
import json
import re
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

MAX_TEXT = 100_000
MAX_PAGES = 40
MAX_OCR_PAGES = 8


class ParseError(ValueError):
    pass


def checked_zip(path):
    archive = zipfile.ZipFile(path)
    infos = archive.infolist()
    if len(infos) > 4000 or sum(i.file_size for i in infos) > 80 * 1024 * 1024:
        archive.close()
        raise ParseError("office_archive_limit")
    if any(i.flag_bits & 1 for i in infos):
        archive.close()
        raise ParseError("encrypted_document")
    return archive


def xml(archive, name):
    raw = archive.read(name)
    if b"<!DOCTYPE" in raw.upper() or b"<!ENTITY" in raw.upper():
        raise ParseError("xml_entities_not_allowed")
    return ET.fromstring(raw)


def decode_text(content):
    if content.startswith((b"\xff\xfe", b"\xfe\xff")):
        return content.decode("utf-16")
    if b"\x00" in content:
        raise ParseError("binary_file_as_text")
    try:
        return content.decode("utf-8-sig")
    except UnicodeDecodeError:
        return content.decode("cp1251")


def packetize(blocks):
    chunks = []
    total = 0
    for block in blocks:
        text = block["text"].strip()
        total += len(text)
        if total > MAX_TEXT:
            raise ParseError("document_text_limit")
        for start in range(0, len(text), 9000):
            part = text[start : start + 9000]
            if not part.strip():
                continue
            chunks.append(
                {
                    "id": f"chunk_{len(chunks) + 1}",
                    "text": part,
                    "locator": block["locator"]
                    + (f"; символы {start + 1}–{min(start + 9000, len(text))}" if len(text) > 9000 else ""),
                    "page": block.get("page"),
                    "method": block.get("method", "text"),
                }
            )
    return chunks


def parse_document(path):
    path = Path(path)
    extension = path.suffix.lower()
    blocks, ocr_pages, issues = [], [], []
    if extension == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(path)
        if reader.is_encrypted:
            raise ParseError("encrypted_document")
        if len(reader.pages) > MAX_PAGES:
            raise ParseError("pdf_page_limit")
        for index, page in enumerate(reader.pages, 1):
            if float(page.mediabox.width) > 14400 or float(page.mediabox.height) > 14400:
                raise ParseError("pdf_page_dimensions_limit")
            if (
                len(page.get_contents().get_data()) > 10 * 1024 * 1024
                if page.get_contents() is not None
                else False
            ):
                raise ParseError("pdf_page_stream_limit")
            text = page.extract_text() or ""
            if len(text.strip()) < 30:
                ocr_pages.append(index)
            else:
                blocks.append({"text": text, "locator": f"Страница {index}", "page": index})
        if len(ocr_pages) > MAX_OCR_PAGES:
            raise ParseError("scanned_pdf_page_limit")
    elif extension in {".docx", ".pptx"}:
        with checked_zip(path) as archive:
            if extension == ".docx":
                tree = xml(archive, "word/document.xml")
                ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
                paragraphs = [
                    "".join(n.text or "" for n in p.findall(".//w:t", ns)) for p in tree.findall(".//w:p", ns)
                ]
                text = "\n".join(paragraphs)
                blocks.append({"text": text, "locator": "Основной текст DOCX, включая ячейки таблиц"})
                if any("media/" in n for n in archive.namelist()):
                    issues.append(
                        "Изображения DOCX не распознаны; текстовые блоки и таблицы прочитаны. Числа только на изображениях недоступны."
                    )
            else:
                slides = sorted(
                    (n for n in archive.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)),
                    key=lambda n: int(re.search(r"slide(\d+)", n)[1]),
                )
                if len(slides) > 60:
                    raise ParseError("presentation_slide_limit")
                ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
                for index, name in enumerate(slides, 1):
                    tree = xml(archive, name)
                    text = "\n".join(n.text or "" for n in tree.findall(".//a:t", ns))
                    blocks.append({"text": text, "locator": f"Слайд {index}", "page": index})
                issues.append(
                    "Прочитаны текстовые блоки и таблицы слайдов; диаграммы и изображения без текстовых подписей не оцифрованы."
                )
    elif extension == ".xlsx":
        from openpyxl import load_workbook

        with checked_zip(path) as archive:
            if any("vbaProject" in n for n in archive.namelist()):
                raise ParseError("macros_not_allowed")
            for name in archive.namelist():
                if name.endswith(".xml"):
                    raw = archive.read(name)
                    if b"<!DOCTYPE" in raw.upper() or b"<!ENTITY" in raw.upper():
                        raise ParseError("xml_entities_not_allowed")
        book = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        try:
            cell_count = 0
            for sheet in book:
                if (sheet.max_row or 0) > 5000 or (sheet.max_column or 0) > 200:
                    raise ParseError("spreadsheet_cell_limit")
                lines = []
                for index, row in enumerate(sheet.iter_rows(values_only=True), 1):
                    cell_count += len(row)
                    if index > 5000 or cell_count > 100_000:
                        raise ParseError("spreadsheet_cell_limit")
                    lines.append(
                        f"Строка {index}: " + " | ".join(str(v) if v is not None else "" for v in row)
                    )
                blocks.append({"text": "\n".join(lines), "locator": "Лист «" + sheet.title + "»"})
        finally:
            book.close()
        issues.append(
            "Формулы и внешние ссылки не выполнялись; прочитаны константы и сохранённые в файле значения формул."
        )
    elif extension in {".txt", ".md", ".csv"}:
        text = decode_text(path.read_bytes())
        if extension == ".csv":
            try:
                dialect = csv.Sniffer().sniff(text[:4096], delimiters=";,\t,")
            except csv.Error:
                dialect = csv.excel
            rows = []
            for index, row in enumerate(csv.reader(io.StringIO(text), dialect), 1):
                if index > 5000 or len(row) > 200:
                    raise ParseError("spreadsheet_cell_limit")
                rows.append(f"Строка {index}: " + " | ".join(row))
            text = "\n".join(rows)
        blocks.append(
            {"text": text, "locator": "Текст файла" if extension != ".csv" else "CSV: строки с заголовками"}
        )
    else:
        raise ParseError("file_type_not_supported")
    chunks = packetize(blocks)
    if not chunks and not ocr_pages:
        raise ParseError("document_has_no_readable_content")
    return {"chunks": chunks, "ocr_pages": ocr_pages, "issues": issues}


def render_page(path, page_number):
    import pypdfium2 as pdfium

    document = pdfium.PdfDocument(path)
    try:
        page = document[page_number - 1]
        width, height = page.get_size()
        scale = min(2.0, 1600 / max(width, height))
        bitmap = page.render(scale=scale)
        try:
            image = bitmap.to_pil()
            output = io.BytesIO()
            image.save(output, format="PNG")
            return output.getvalue()
        finally:
            bitmap.close()
            page.close()
    finally:
        document.close()


if __name__ == "__main__":
    # Called in a separate bounded process; stdout contains JSON, never diagnostics.
    try:
        print(json.dumps(parse_document(Path(sys.argv[1])), ensure_ascii=False))
    except Exception as error:
        print(json.dumps({"error": str(error) if isinstance(error, ParseError) else "file_parse_failed"}))
        sys.exit(2)
